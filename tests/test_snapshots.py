from __future__ import annotations

import json
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from pydantic import ValidationError

from forge.ingest.adapters import (
    LocalFileSourceAdapter,
    SourceArtifact,
    SourceRef,
)
from forge.ingest.document import (
    add_product_context,
    add_supplemental_answers,
    ingest_document,
    ingest_snapshot,
)
from forge.ingest.batching import batch_document, plan_fingerprint
from forge.ingest.models import ProductContextTerm, SupplementalAnswer
from forge.ingest.parsers import LegacyDocumentParser
from forge.ingest.snapshot import SnapshotCache, build_snapshot, normalized_hash
from forge.rubric.loader import load_rubric


def test_local_adapter_reads_bytes_once_and_retains_source_metadata(
    tmp_path, monkeypatch
):
    path = tmp_path / "product brief.md"
    path.write_bytes(b"# Goal\nShip safely.\n")
    reads = 0
    original_read_bytes = Path.read_bytes

    def counted_read_bytes(candidate):
        nonlocal reads
        reads += 1
        return original_read_bytes(candidate)

    monkeypatch.setattr(Path, "read_bytes", counted_read_bytes)

    artifact = LocalFileSourceAdapter().acquire(
        SourceRef.local_file(path).model_copy(
            update={"origin_metadata": {"workspace": "alpha"}}
        )
    )
    parsed = LegacyDocumentParser().parse(artifact)

    assert reads == 1
    assert artifact.content == b"# Goal\nShip safely.\n"
    assert artifact.display_name == "product brief.md"
    assert artifact.source_type == "md"
    assert artifact.media_type == "text/markdown"
    assert artifact.origin_locator == str(path.resolve())
    assert artifact.origin_metadata == {
        "workspace": "alpha",
        "resolved_path": str(path.resolve()),
    }
    assert parsed.document.blocks[0].text == "# Goal\nShip safely."


def test_source_ref_allows_future_origins_but_local_adapter_does_not_fetch_them():
    source = SourceRef(
        origin="external",
        locator="https://example.invalid/prd.pdf",
        display_name="prd.pdf",
        source_type="pdf",
    )

    with pytest.raises(ValueError, match="not implemented"):
        LocalFileSourceAdapter().acquire(source)


def test_same_content_at_different_paths_has_same_snapshot_identity(tmp_path):
    first = tmp_path / "first.md"
    second = tmp_path / "nested" / "second.md"
    second.parent.mkdir()
    first.write_bytes(b"# Requirement\nRetry failed requests.\n")
    second.write_bytes(first.read_bytes())

    first_snapshot = ingest_snapshot(first, SnapshotCache())
    second_snapshot = ingest_snapshot(second, SnapshotCache())

    assert first_snapshot.snapshot_id == second_snapshot.snapshot_id
    assert first_snapshot.normalized_hash == second_snapshot.normalized_hash
    assert first_snapshot.source.origin_locator != second_snapshot.source.origin_locator
    assert first_snapshot.document.source_path != second_snapshot.document.source_path


def test_parser_fingerprint_changes_snapshot_identity(tmp_path):
    path = tmp_path / "prd.txt"
    path.write_text("One stable requirement.")
    artifact = LocalFileSourceAdapter().acquire(SourceRef.local_file(path))
    parsed = LegacyDocumentParser().parse(artifact)

    original = build_snapshot(artifact, parsed, LegacyDocumentParser.fingerprint)
    changed = build_snapshot(artifact, parsed, "forge.legacy-document-parser/2.0")

    assert original.normalized_hash == changed.normalized_hash
    assert original.snapshot_id != changed.snapshot_id


def test_snapshot_metadata_does_not_affect_normalized_hash(tmp_path):
    path = tmp_path / "prd.txt"
    path.write_text("One stable requirement.")
    snapshot = ingest_snapshot(path, SnapshotCache())
    changed_metadata = snapshot.document.model_copy(
        update={
            "snapshot_id": "different",
            "source_sha256": "different",
            "parser_fingerprint": "different",
            "normalized_hash": "different",
            "normalized_schema_version": "different",
        }
    )

    assert normalized_hash(
        changed_metadata, snapshot.canonical_nodes
    ) == normalized_hash(snapshot.document, snapshot.canonical_nodes)


def test_v2_plan_fingerprint_changes_with_parser_and_normalized_snapshot(tmp_path):
    path = tmp_path / "prd.txt"
    path.write_text("One stable requirement.")
    artifact = LocalFileSourceAdapter().acquire(SourceRef.local_file(path))
    parsed = LegacyDocumentParser().parse(artifact)
    original = build_snapshot(artifact, parsed, LegacyDocumentParser.fingerprint)
    changed_parser = build_snapshot(artifact, parsed, "test-parser/2")
    changed_document = parsed.document.model_copy(
        update={
            "blocks": [
                parsed.document.blocks[0].model_copy(
                    update={"text": "A different normalized requirement."}
                )
            ]
        }
    )
    changed_normalized = build_snapshot(
        artifact,
        parsed.model_copy(update={"document": changed_document}),
        LegacyDocumentParser.fingerprint,
    )
    rubric = load_rubric("prd")

    original_fingerprint = plan_fingerprint(
        batch_document(original.document), rubric.id, rubric.version
    )
    assert original_fingerprint != plan_fingerprint(
        batch_document(changed_parser.document), rubric.id, rubric.version
    )
    assert original_fingerprint != plan_fingerprint(
        batch_document(changed_normalized.document), rubric.id, rubric.version
    )


def test_snapshot_json_round_trip_uses_base64_for_arbitrary_source_bytes(tmp_path):
    path = tmp_path / "prd.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "A binary source snapshot.")
    pdf.save(path)
    pdf.close()

    snapshot = ingest_snapshot(path, SnapshotCache())
    payload = snapshot.model_dump_json()
    restored = type(snapshot).model_validate_json(payload)

    assert restored == snapshot
    assert json.loads(payload)["source"]["content"] != snapshot.source.content


def test_snapshot_is_frozen_and_context_or_answers_only_change_projections(tmp_path):
    path = tmp_path / "prd.md"
    path.write_text("# Requirement\nExport invoices.")
    snapshot = ingest_snapshot(path, SnapshotCache())

    with pytest.raises(ValidationError, match="frozen"):
        snapshot.snapshot_id = "0" * 64

    with_context = add_product_context(
        snapshot.document,
        [ProductContextTerm(term="invoice", meaning="customer billing record")],
    )
    with_answer = add_supplemental_answers(
        snapshot.document,
        [SupplementalAnswer(criterion_id="scope", answer="Admins are in scope.")],
    )

    assert snapshot.document.product_context == []
    assert all(block.provenance == "document" for block in snapshot.document.blocks)
    assert with_context.product_context
    assert with_answer.blocks[-1].provenance == "supplemental_answer"


def test_snapshot_cache_uses_only_the_exact_key(tmp_path):
    path = tmp_path / "prd.txt"
    path.write_text("Cache this source.")
    cache = SnapshotCache()

    first = ingest_snapshot(path, cache)
    second = ingest_snapshot(path, cache)
    path.write_text("Cache this changed source.")
    changed = ingest_snapshot(path, cache)

    assert second is first
    assert changed.snapshot_id != first.snapshot_id


def test_compatibility_projection_preserves_legacy_shapes_for_all_formats(tmp_path):
    md_path = tmp_path / "brief.md"
    txt_path = tmp_path / "brief.txt"
    md_path.write_bytes(b"# Problem\r\nManual work.\r\n")
    txt_path.write_text("Plain requirement.")

    docx_path = tmp_path / "brief.docx"
    docx = Document()
    docx.sections[0].header.paragraphs[0].text = "Owner: Product"
    docx.add_heading("Scope", level=1)
    docx.add_paragraph("Include administrators.")
    table = docx.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "State"
    table.cell(0, 1).text = "Ready"
    docx.save(docx_path)

    pdf_path = tmp_path / "brief.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "PDF requirement.")
    page.draw_rect(pymupdf.Rect(72, 100, 120, 140))
    pdf.save(pdf_path)
    pdf.close()

    assert ingest_document(md_path).model_dump() == {
        "source_path": str(md_path.resolve()),
        "source_type": "md",
        "blocks": [
            {
                "id": "text-1",
                "text": "# Problem\nManual work.",
                "page": None,
                "section": None,
                "provenance": "document",
                "criterion_id": None,
                "parent_id": None,
                "start_char": None,
                "end_char": None,
            }
        ],
        "visual_assets": [],
        "product_context": [],
        "snapshot_id": ingest_snapshot(md_path).snapshot_id,
        "source_sha256": ingest_snapshot(md_path).source.sha256,
        "parser_fingerprint": ingest_snapshot(md_path).parser_fingerprint,
        "normalized_hash": ingest_snapshot(md_path).normalized_hash,
        "normalized_schema_version": ingest_snapshot(
            md_path
        ).normalized_schema_version,
    }
    assert [block.text for block in ingest_document(txt_path).blocks] == [
        "Plain requirement."
    ]
    assert [block.model_dump() for block in ingest_document(docx_path).blocks] == [
        {
            "id": "paragraph-1",
            "text": "Scope",
            "page": None,
            "section": "Scope",
            "provenance": "document",
            "criterion_id": None,
            "parent_id": None,
            "start_char": None,
            "end_char": None,
        },
        {
            "id": "paragraph-2",
            "text": "Include administrators.",
            "page": None,
            "section": "Scope",
            "provenance": "document",
            "criterion_id": None,
            "parent_id": None,
            "start_char": None,
            "end_char": None,
        },
        {
            "id": "table-1-row-1",
            "text": "State | Ready",
            "page": None,
            "section": "Scope",
            "provenance": "document",
            "criterion_id": None,
            "parent_id": None,
            "start_char": None,
            "end_char": None,
        },
        {
            "id": "section-1-header-1",
            "text": "Owner: Product",
            "page": None,
            "section": "Header",
            "provenance": "document",
            "criterion_id": None,
            "parent_id": None,
            "start_char": None,
            "end_char": None,
        },
    ]
    pdf_document = ingest_document(pdf_path)
    assert [block.model_dump() for block in pdf_document.blocks] == [
        {
            "id": "page-1",
            "text": "PDF requirement.",
            "page": 1,
            "section": None,
            "provenance": "document",
            "criterion_id": None,
            "parent_id": None,
            "start_char": None,
            "end_char": None,
        }
    ]
    assert [asset.model_dump() for asset in pdf_document.visual_assets] == [
        {
            "id": "page-1-visual",
            "media_type": "application/pdf-page",
            "page": 1,
            "section": None,
            "locator": None,
        }
    ]


def test_canonical_nodes_capture_known_legacy_structure(tmp_path):
    path = tmp_path / "structured.docx"
    document = Document()
    document.sections[0].footer.paragraphs[0].text = "Confidential"
    document.add_heading("Requirements", level=1)
    document.add_paragraph("Retry on timeout.")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Timeout | Retry"
    document.save(path)

    nodes = ingest_snapshot(path, SnapshotCache()).canonical_nodes

    assert [node.kind for node in nodes] == [
        "heading",
        "paragraph",
        "table_row",
        "footer",
    ]
    assert [node.node_id for node in nodes] == [
        "paragraph-1",
        "paragraph-2",
        "table-1-row-1",
        "section-1-footer-1",
    ]
    assert [node.order for node in nodes] == list(range(len(nodes)))
    assert nodes[1].heading_path == ("Requirements",)
    assert nodes[2].table.table_id == "table-1"
