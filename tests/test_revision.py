from __future__ import annotations

from docx import Document
import pytest

from forge.ingest.document import ingest_document
from forge.ingest.models import SupplementalAnswer
from forge.revise import materialize_prd_revision
from forge.revise import (
    approve_revision_plan,
    materialize_integrated_prd_revision,
    preview_integrated_revision,
)


def test_materializes_answers_into_new_markdown_revision(tmp_path):
    source = tmp_path / "prd.md"
    output = tmp_path / "prd-revised.md"
    source.write_text("# Original PRD\n\nUsers cannot export invoices.\n")
    answers = [
        SupplementalAnswer(
            criterion_id="problem_statement",
            answer="Finance administrators are affected, based on 42 tickets.",
        ),
        SupplementalAnswer(
            criterion_id="success_metrics",
            answer="Reduce export failures from 8% to 1% within 30 days.",
        ),
    ]

    result = materialize_prd_revision(source, output, answers)

    assert source.read_text() == "# Original PRD\n\nUsers cannot export invoices.\n"
    assert result.supplemental_answer_count == 2
    revised = output.read_text()
    assert "# Forge Clarifications" in revised
    assert "## Problem Statement" in revised
    assert answers[0].answer in revised
    assert ingest_document(output).locate_quote(answers[1].answer) is not None


def test_materializes_answers_into_new_docx_revision(tmp_path):
    source = tmp_path / "prd.docx"
    output = tmp_path / "prd-revised.docx"
    document = Document()
    document.add_heading("Original PRD", level=1)
    document.add_paragraph("Users cannot export invoices.")
    document.save(source)
    answer = SupplementalAnswer(
        criterion_id="problem_statement",
        answer="Finance administrators are affected, based on 42 tickets.",
    )

    materialize_prd_revision(source, output, [answer])

    assert answer.answer in [paragraph.text for paragraph in Document(output).paragraphs]
    assert ingest_document(output).locate_quote(answer.answer) is not None


def test_revision_never_overwrites_source_or_existing_output(tmp_path):
    source = tmp_path / "prd.txt"
    source.write_text("Original")
    answer = SupplementalAnswer(criterion_id="non_goals", answer="Mobile is excluded.")

    with pytest.raises(ValueError, match="must differ"):
        materialize_prd_revision(source, source, [answer])

    output = tmp_path / "existing.txt"
    output.write_text("Do not replace")
    with pytest.raises(ValueError, match="already exists"):
        materialize_prd_revision(source, output, [answer])


def test_previews_approves_and_integrates_markdown_revision(tmp_path):
    source = tmp_path / "prd.md"
    output = tmp_path / "prd-revised.md"
    source.write_text("# PRD\n\n## Success Metrics\n\nDAU is monitored.\n")
    answer = SupplementalAnswer(
        criterion_id="success_metrics",
        answer="Primary metric: repeat usage increases from 20% to 30% in 90 days.",
    )

    plan = preview_integrated_revision(source, [answer])
    assert not output.exists()
    assert plan.edits[0].target_section == "Success Metrics"
    assert plan.edits[0].existing_excerpt == "DAU is monitored."
    assert plan.edits[0].conflicts[0].kind == "existing_section_content"
    approved = approve_revision_plan(
        plan, actions={plan.edits[0].edit_id: "integrate"}
    )
    result = materialize_integrated_prd_revision(source, output, approved)

    assert source.read_text() == "# PRD\n\n## Success Metrics\n\nDAU is monitored.\n"
    revised = output.read_text()
    assert answer.answer in revised
    assert revised.index(answer.answer) < revised.index("# Forge Revision Audit")
    assert approved.approval_digest in revised
    assert result.output_path == str(output.resolve())


def test_integrated_revision_rejects_stale_source_and_modified_plan(tmp_path):
    source = tmp_path / "prd.md"
    source.write_text("# PRD\n\n## Problem\n\nCurrent text.\n")
    answer = SupplementalAnswer(
        criterion_id="problem_statement",
        answer="Mobile viewers cannot find short-form stories.",
    )
    plan = preview_integrated_revision(source, [answer])
    approved = approve_revision_plan(
        plan, actions={plan.edits[0].edit_id: "integrate"}
    )
    source.write_text("# PRD\n\n## Problem\n\nChanged text.\n")

    with pytest.raises(ValueError, match="source changed"):
        materialize_integrated_prd_revision(
            source, tmp_path / "revised.md", approved
        )

    source.write_text("# PRD\n\n## Problem\n\nCurrent text.\n")
    approved.edits[0].edit.answer = "Tampered"
    with pytest.raises(ValueError, match="modified after approval"):
        materialize_integrated_prd_revision(
            source, tmp_path / "revised.md", approved
        )


def test_missing_revision_section_requires_explicit_audit_resolution(tmp_path):
    source = tmp_path / "prd.md"
    output = tmp_path / "prd-revised.md"
    source.write_text("# PRD\n\nContent.\n")
    answer = SupplementalAnswer(
        criterion_id="risk_compliance",
        answer="No customer PII is stored.",
    )
    plan = preview_integrated_revision(source, [answer])
    edit = plan.edits[0]
    assert edit.target_section is None
    assert edit.conflicts[0].kind == "missing_section"

    with pytest.raises(ValueError, match="has no integration target"):
        approve_revision_plan(plan, actions={edit.edit_id: "integrate"})

    approved = approve_revision_plan(plan, actions={edit.edit_id: "audit_only"})
    materialize_integrated_prd_revision(source, output, approved)
    assert answer.answer in output.read_text()
