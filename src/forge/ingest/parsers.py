from __future__ import annotations

from io import BytesIO, TextIOWrapper
from typing import Protocol

import pymupdf
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.table import Table
from docx.text.paragraph import Paragraph
from pydantic import BaseModel, ConfigDict

from forge.ingest.adapters import SourceArtifact
from forge.ingest.models import NormalizedDocument, SourceBlock, VisualAsset
from forge.ingest.nodes import (
    AssetNodeMetadata,
    CanonicalNode,
    NodeProvenance,
    TableNodeMetadata,
)


LEGACY_PARSER_FINGERPRINT = "forge.legacy-document-parser/1.0"


class ParsedDocument(BaseModel):
    model_config = ConfigDict(frozen=True)

    document: NormalizedDocument
    canonical_nodes: tuple[CanonicalNode, ...]


class DocumentParser(Protocol):
    fingerprint: str

    def parse(self, artifact: SourceArtifact) -> ParsedDocument: ...


class LegacyDocumentParser:
    """The original Forge normalization behavior, pinned to an explicit version."""

    fingerprint = LEGACY_PARSER_FINGERPRINT

    def parse(self, artifact: SourceArtifact) -> ParsedDocument:
        if artifact.source_type == "pdf":
            blocks, visual_assets, nodes = _pdf_content(artifact)
        elif artifact.source_type == "docx":
            blocks, visual_assets, nodes = _docx_content(artifact)
        elif artifact.source_type in {"md", "txt"}:
            blocks, nodes = _text_content(artifact)
            visual_assets = []
        else:
            raise ValueError(f"unsupported source type {artifact.source_type!r}")

        source_path = artifact.origin_locator or artifact.display_name
        if not blocks and not visual_assets:
            raise ValueError(f"document contains no extractable text: {source_path}")
        return ParsedDocument(
            document=NormalizedDocument(
                source_path=source_path,
                source_type=artifact.source_type,
                blocks=blocks,
                visual_assets=visual_assets,
            ),
            canonical_nodes=tuple(nodes),
        )


def _pdf_content(
    artifact: SourceArtifact,
) -> tuple[list[SourceBlock], list[VisualAsset], list[CanonicalNode]]:
    blocks: list[SourceBlock] = []
    visual_assets: list[VisualAsset] = []
    nodes: list[CanonicalNode] = []
    with pymupdf.open(stream=artifact.content, filetype="pdf") as pdf:
        for page_number, page in enumerate(pdf, start=1):
            text = page.get_text("text").strip()
            if text:
                block = SourceBlock(
                    id=f"page-{page_number}", text=text, page=page_number
                )
                blocks.append(block)
                nodes.append(
                    CanonicalNode(
                        node_id=block.id,
                        order=len(nodes),
                        kind="page",
                        content=text,
                        page=page_number,
                        provenance=NodeProvenance(
                            source_block_id=block.id, parser=LEGACY_PARSER_FINGERPRINT
                        ),
                    )
                )
            if page.get_images(full=True) or page.get_drawings():
                asset = VisualAsset(
                    id=f"page-{page_number}-visual",
                    media_type="application/pdf-page",
                    page=page_number,
                )
                visual_assets.append(asset)
                nodes.append(_visual_node(asset, len(nodes)))
    return blocks, visual_assets, nodes


def _docx_content(
    artifact: SourceArtifact,
) -> tuple[list[SourceBlock], list[VisualAsset], list[CanonicalNode]]:
    document = Document(BytesIO(artifact.content))
    blocks: list[SourceBlock] = []
    nodes: list[CanonicalNode] = []
    section: str | None = None
    heading_path: list[str] = []
    index = 0
    table_number = 0

    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            index += 1
            kind = "paragraph"
            if item.style and item.style.name.startswith("Heading"):
                section = text
                kind = "heading"
                try:
                    level = int(item.style.name.removeprefix("Heading "))
                except ValueError:
                    level = 1
                heading_path[level - 1 :] = [text]
            block = SourceBlock(id=f"paragraph-{index}", text=text, section=section)
            blocks.append(block)
            nodes.append(
                CanonicalNode(
                    node_id=block.id,
                    order=len(nodes),
                    kind=kind,
                    content=text,
                    heading_path=tuple(heading_path),
                    provenance=NodeProvenance(
                        source_block_id=block.id, parser=LEGACY_PARSER_FINGERPRINT
                    ),
                )
            )
        elif isinstance(item, Table):
            table_number += 1
            for row_number, row in enumerate(item.rows, start=1):
                cells = [cell.text.strip() for cell in row.cells]
                text = " | ".join(cell for cell in cells if cell)
                if text:
                    block = SourceBlock(
                        id=f"table-{table_number}-row-{row_number}",
                        text=text,
                        section=section,
                    )
                    blocks.append(block)
                    nodes.append(
                        CanonicalNode(
                            node_id=block.id,
                            order=len(nodes),
                            kind="table_row",
                            content=text,
                            heading_path=tuple(heading_path),
                            table=TableNodeMetadata(
                                table_id=f"table-{table_number}", row=row_number
                            ),
                            provenance=NodeProvenance(
                                source_block_id=block.id,
                                parser=LEGACY_PARSER_FINGERPRINT,
                            ),
                        )
                    )

    seen_parts: set[str] = set()
    for section_number, doc_section in enumerate(document.sections, start=1):
        for kind, container in (
            ("header", doc_section.header),
            ("footer", doc_section.footer),
        ):
            part_name = str(container.part.partname)
            if part_name in seen_parts:
                continue
            seen_parts.add(part_name)
            for paragraph_number, paragraph in enumerate(
                container.paragraphs, start=1
            ):
                text = paragraph.text.strip()
                if text:
                    block = SourceBlock(
                        id=f"section-{section_number}-{kind}-{paragraph_number}",
                        text=text,
                        section=kind.title(),
                    )
                    blocks.append(block)
                    nodes.append(
                        CanonicalNode(
                            node_id=block.id,
                            order=len(nodes),
                            kind=kind,
                            content=text,
                            heading_path=(kind.title(),),
                            provenance=NodeProvenance(
                                source_block_id=block.id,
                                parser=LEGACY_PARSER_FINGERPRINT,
                            ),
                        )
                    )

    image_relationships = sorted(
        (
            relationship
            for relationship in document.part.rels.values()
            if relationship.reltype == RELATIONSHIP_TYPE.IMAGE
        ),
        key=lambda relationship: relationship.rId,
    )
    visual_assets = [
        VisualAsset(
            id=f"embedded-image-{index}",
            media_type=relationship.target_part.content_type,
            locator=relationship.rId,
        )
        for index, relationship in enumerate(image_relationships, start=1)
    ]
    first_visual_order = len(nodes)
    nodes.extend(
        _visual_node(asset, first_visual_order + index)
        for index, asset in enumerate(visual_assets)
    )
    return blocks, visual_assets, nodes


def _text_content(
    artifact: SourceArtifact,
) -> tuple[list[SourceBlock], list[CanonicalNode]]:
    with TextIOWrapper(BytesIO(artifact.content), encoding="utf-8") as stream:
        text = stream.read().strip()
    if not text:
        return [], []
    block = SourceBlock(id="text-1", text=text)
    return [block], [
        CanonicalNode(
            node_id=block.id,
            order=0,
            kind="paragraph",
            content=text,
            provenance=NodeProvenance(
                source_block_id=block.id, parser=LEGACY_PARSER_FINGERPRINT
            ),
        )
    ]


def _visual_node(asset: VisualAsset, order: int) -> CanonicalNode:
    return CanonicalNode(
        node_id=asset.id,
        order=order,
        kind="visual",
        page=asset.page,
        asset=AssetNodeMetadata(
            asset_id=asset.id,
            media_type=asset.media_type,
            locator=asset.locator,
        ),
        provenance=NodeProvenance(parser=LEGACY_PARSER_FINGERPRINT),
    )
