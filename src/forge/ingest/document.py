from __future__ import annotations

from pathlib import Path

import pymupdf
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.table import Table
from docx.text.paragraph import Paragraph

from forge.ingest.models import (
    NormalizedDocument,
    ProductContextTerm,
    SourceBlock,
    SupplementalAnswer,
    VisualAsset,
)


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".txt"}


def ingest_document(source: str | Path) -> NormalizedDocument:
    path = Path(source).expanduser().resolve()
    if not path.exists() or not path.is_file():
        raise FileNotFoundError(f"document not found: {path}")
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"unsupported document type {path.suffix!r}; use {supported}")

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        blocks, visual_assets = _pdf_content(path)
    elif suffix == ".docx":
        blocks, visual_assets = _docx_content(path)
    else:
        blocks = _text_blocks(path)
        visual_assets = []

    if not blocks and not visual_assets:
        raise ValueError(f"document contains no extractable text: {path}")
    return NormalizedDocument(
        source_path=str(path),
        source_type=suffix.removeprefix("."),
        blocks=blocks,
        visual_assets=visual_assets,
    )


def add_supplemental_answers(
    document: NormalizedDocument, answers: list[SupplementalAnswer]
) -> NormalizedDocument:
    blocks = list(document.blocks)
    blocks.extend(
        SourceBlock(
            id=f"supplemental-answer-{index}",
            text=answer.answer,
            provenance="supplemental_answer",
            criterion_id=answer.criterion_id,
        )
        for index, answer in enumerate(answers, start=1)
    )
    return document.model_copy(update={"blocks": blocks})


def add_product_context(
    document: NormalizedDocument, terms: list[ProductContextTerm]
) -> NormalizedDocument:
    """Attach non-evidence terminology context for extraction disambiguation."""
    return document.model_copy(update={"product_context": terms})


def _pdf_content(path: Path) -> tuple[list[SourceBlock], list[VisualAsset]]:
    blocks: list[SourceBlock] = []
    visual_assets: list[VisualAsset] = []
    with pymupdf.open(path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            text = page.get_text("text").strip()
            if text:
                blocks.append(
                    SourceBlock(id=f"page-{page_number}", text=text, page=page_number)
                )
            if page.get_images(full=True) or page.get_drawings():
                visual_assets.append(
                    VisualAsset(
                        id=f"page-{page_number}-visual",
                        media_type="application/pdf-page",
                        page=page_number,
                    )
                )
    return blocks, visual_assets


def _docx_content(path: Path) -> tuple[list[SourceBlock], list[VisualAsset]]:
    document = Document(path)
    blocks: list[SourceBlock] = []
    section: str | None = None
    index = 0

    table_number = 0
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            index += 1
            if item.style and item.style.name.startswith("Heading"):
                section = text
            blocks.append(
                SourceBlock(id=f"paragraph-{index}", text=text, section=section)
            )
        elif isinstance(item, Table):
            table_number += 1
            for row_number, row in enumerate(item.rows, start=1):
                cells = [cell.text.strip() for cell in row.cells]
                text = " | ".join(cell for cell in cells if cell)
                if text:
                    blocks.append(
                        SourceBlock(
                            id=f"table-{table_number}-row-{row_number}",
                            text=text,
                            section=section,
                        )
                    )

    seen_parts: set[str] = set()
    for section_number, doc_section in enumerate(document.sections, start=1):
        for kind, container in (
            ("header", doc_section.header),
            ("footer", doc_section.footer),
        ):
            part_name = str(container.part.partname)
            if part_name in seen_parts:
                continue
            seen_parts.add(part_name)
            for paragraph_number, paragraph in enumerate(
                container.paragraphs, start=1
            ):
                text = paragraph.text.strip()
                if text:
                    blocks.append(
                        SourceBlock(
                            id=(
                                f"section-{section_number}-{kind}-"
                                f"{paragraph_number}"
                            ),
                            text=text,
                            section=kind.title(),
                        )
                    )
    image_relationships = sorted(
        (
            relationship
            for relationship in document.part.rels.values()
            if relationship.reltype == RELATIONSHIP_TYPE.IMAGE
        ),
        key=lambda relationship: relationship.rId,
    )
    visual_assets = [
        VisualAsset(
            id=f"embedded-image-{index}",
            media_type=relationship.target_part.content_type,
            locator=relationship.rId,
        )
        for index, relationship in enumerate(image_relationships, start=1)
    ]
    return blocks, visual_assets


def _text_blocks(path: Path) -> list[SourceBlock]:
    text = path.read_text(encoding="utf-8").strip()
    return [SourceBlock(id="text-1", text=text)] if text else []
