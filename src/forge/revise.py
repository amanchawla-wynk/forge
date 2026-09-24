"""Explicitly materialize conversational answers into a new PRD revision."""

from __future__ import annotations

from collections import defaultdict
import hashlib
import json
from pathlib import Path
import re
from typing import Literal
import uuid

from docx import Document
from docx.text.paragraph import Paragraph
from pydantic import BaseModel, Field

from forge.ingest.document import SUPPORTED_EXTENSIONS
from forge.ingest.models import SupplementalAnswer


class RevisionResult(BaseModel):
    source_path: str
    output_path: str
    supplemental_answer_count: int
    note: str
    mode: Literal["appendix", "integrated"] = "appendix"
    plan_digest: str | None = None
    final_assessment_required: bool = False


class RevisionConflict(BaseModel):
    conflict_id: str
    edit_id: str
    kind: Literal[
        "missing_section", "duplicate_text", "existing_section_content"
    ]
    message: str
    allowed_resolutions: list[str]


class RevisionEdit(BaseModel):
    edit_id: str
    criterion_id: str
    answer: str
    target_section: str | None = None
    existing_excerpt: str | None = None
    conflicts: list[RevisionConflict] = Field(default_factory=list)


class RevisionPlan(BaseModel):
    version: Literal["1.0"] = "1.0"
    source_path: str
    source_sha256: str
    source_type: str
    edits: list[RevisionEdit]
    plan_digest: str


class ApprovedRevisionEdit(BaseModel):
    edit: RevisionEdit
    action: Literal["integrate", "audit_only", "skip"]


class ApprovedRevisionPlan(BaseModel):
    plan: RevisionPlan
    edits: list[ApprovedRevisionEdit]
    approval_digest: str


_SECTION_ALIASES: dict[str, tuple[str, ...]] = {
    "problem_statement": ("problem", "introduction", "background", "opportunity"),
    "success_metrics": ("success", "metric", "measurement"),
    "non_goals": ("non-goal", "non goal", "out of scope", "scope"),
    "functional_requirements": ("requirement", "experience", "flow"),
    "acceptance_criteria": ("acceptance", "success criterion"),
    "edge_cases_and_states": ("edge case", "error", "loading", "state"),
    "instrumentation": ("analytics", "instrumentation", "observability"),
    "dependencies": ("dependencies", "partner", "constraint"),
    "risk_compliance": ("risk", "privacy", "security", "compliance"),
    "rollout": ("rollout", "launch", "poc", "experiment"),
    "open_questions_owned": ("open question", "open point", "doubt"),
    "alternatives_considered": ("alternative", "option"),
    "assumptions_validation": ("assumption", "validation", "poc"),
    "operational_readiness": ("operation", "support", "recovery"),
    "document_governance": ("owner", "governance", "status"),
}


def materialize_prd_revision(
    source_path: str | Path,
    output_path: str | Path,
    supplemental_answers: list[SupplementalAnswer],
) -> RevisionResult:
    """Write answers into a new document while preserving the original."""
    source = Path(source_path).expanduser().resolve()
    output = Path(output_path).expanduser().resolve()
    if not source.exists() or not source.is_file():
        raise FileNotFoundError(f"document not found: {source}")
    if not supplemental_answers:
        raise ValueError("at least one supplemental answer is required")
    if source == output:
        raise ValueError("output_path must differ from source_path")
    if output.exists():
        raise ValueError(f"output already exists: {output}")
    if not output.parent.exists():
        raise FileNotFoundError(f"output directory not found: {output.parent}")

    suffix = source.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"unsupported document type {suffix!r}; use {supported}")
    if suffix == ".pdf":
        raise ValueError(
            "PDF revisions are not written in place; use the editable DOCX, "
            "Markdown, or text source"
        )
    if output.suffix.lower() != suffix:
        raise ValueError("output_path must use the same file type as source_path")

    grouped: dict[str, list[str]] = defaultdict(list)
    for answer in supplemental_answers:
        grouped[answer.criterion_id].append(answer.answer)

    if suffix == ".docx":
        _write_docx_revision(source, output, grouped)
    else:
        _write_text_revision(source, output, grouped, markdown=suffix == ".md")

    return RevisionResult(
        source_path=str(source),
        output_path=str(output),
        supplemental_answer_count=len(supplemental_answers),
        note=(
            "Created a new revision and preserved the original. Reassess the "
            "output without supplemental answers to score the materialized text."
        ),
    )


def preview_integrated_revision(
    source_path: str | Path,
    supplemental_answers: list[SupplementalAnswer],
    *,
    section_overrides: dict[str, str] | None = None,
) -> RevisionPlan:
    source = Path(source_path).expanduser().resolve()
    _validate_source(source, supplemental_answers)
    sections = _source_sections(source)
    source_text = _source_text(source)
    overrides = section_overrides or {}
    edits: list[RevisionEdit] = []
    for answer in supplemental_answers:
        edit_id = uuid.uuid4().hex
        target = overrides.get(answer.criterion_id) or _best_section(
            answer.criterion_id, sections
        )
        existing_excerpt = _section_excerpt(source, target) if target else None
        conflicts: list[RevisionConflict] = []
        if target is None:
            conflicts.append(
                RevisionConflict(
                    conflict_id=uuid.uuid4().hex,
                    edit_id=edit_id,
                    kind="missing_section",
                    message=(
                        f"No reliable section was found for {answer.criterion_id}."
                    ),
                    allowed_resolutions=["audit_only", "skip"],
                )
            )
        if answer.answer.casefold() in source_text.casefold():
            conflicts.append(
                RevisionConflict(
                    conflict_id=uuid.uuid4().hex,
                    edit_id=edit_id,
                    kind="duplicate_text",
                    message="The approved answer already appears in the source.",
                    allowed_resolutions=["skip", "integrate"],
                )
            )
        elif existing_excerpt:
            conflicts.append(
                RevisionConflict(
                    conflict_id=uuid.uuid4().hex,
                    edit_id=edit_id,
                    kind="existing_section_content",
                    message=(
                        "The target section already contains content. Review the "
                        "existing excerpt and explicitly integrate, append only, "
                        "or skip; Forge does not claim semantic conflict detection "
                        "is exhaustive."
                    ),
                    allowed_resolutions=["integrate", "audit_only", "skip"],
                )
            )
        edits.append(
            RevisionEdit(
                edit_id=edit_id,
                criterion_id=answer.criterion_id,
                answer=answer.answer,
                target_section=target,
                existing_excerpt=existing_excerpt,
                conflicts=conflicts,
            )
        )
    draft = {
        "version": "1.0",
        "source_path": str(source),
        "source_sha256": _file_sha256(source),
        "source_type": source.suffix.lower().removeprefix("."),
        "edits": [edit.model_dump() for edit in edits],
    }
    return RevisionPlan(**draft, plan_digest=_canonical_digest(draft))


def approve_revision_plan(
    plan: RevisionPlan,
    *,
    actions: dict[str, Literal["integrate", "audit_only", "skip"]],
) -> ApprovedRevisionPlan:
    _verify_plan_digest(plan)
    approved: list[ApprovedRevisionEdit] = []
    for edit in plan.edits:
        action = actions.get(edit.edit_id)
        if action is None:
            raise ValueError(f"revision edit {edit.edit_id} has not been approved")
        if action == "integrate" and edit.target_section is None:
            raise ValueError(f"edit {edit.edit_id} has no integration target")
        allowed = {
            resolution
            for conflict in edit.conflicts
            for resolution in conflict.allowed_resolutions
        }
        if edit.conflicts and action not in allowed:
            raise ValueError(
                f"action {action!r} does not resolve conflicts for edit {edit.edit_id}"
            )
        approved.append(ApprovedRevisionEdit(edit=edit, action=action))
    payload = {
        "plan_digest": plan.plan_digest,
        "edits": [item.model_dump() for item in approved],
    }
    return ApprovedRevisionPlan(
        plan=plan,
        edits=approved,
        approval_digest=_canonical_digest(payload),
    )


def materialize_integrated_prd_revision(
    source_path: str | Path,
    output_path: str | Path,
    approved_plan: ApprovedRevisionPlan,
) -> RevisionResult:
    source = Path(source_path).expanduser().resolve()
    output = Path(output_path).expanduser().resolve()
    answers = [
        SupplementalAnswer(
            criterion_id=item.edit.criterion_id,
            answer=item.edit.answer,
        )
        for item in approved_plan.edits
        if item.action != "skip"
    ]
    _validate_paths(source, output, answers)
    approval_payload = {
        "plan_digest": approved_plan.plan.plan_digest,
        "edits": [item.model_dump() for item in approved_plan.edits],
    }
    if _canonical_digest(approval_payload) != approved_plan.approval_digest:
        raise ValueError("approved revision plan was modified after approval")
    _verify_plan_digest(approved_plan.plan)
    if str(source) != approved_plan.plan.source_path:
        raise ValueError("approved revision plan belongs to a different source")
    if _file_sha256(source) != approved_plan.plan.source_sha256:
        raise ValueError("source changed after revision preview")

    integrated = [item for item in approved_plan.edits if item.action == "integrate"]
    if source.suffix.lower() == ".docx":
        _write_integrated_docx(source, output, integrated, approved_plan)
    elif source.suffix.lower() == ".md":
        _write_integrated_markdown(source, output, integrated, approved_plan)
    else:
        _write_integrated_text(source, output, integrated, approved_plan)
    return RevisionResult(
        source_path=str(source),
        output_path=str(output),
        supplemental_answer_count=len(answers),
        note=(
            "Created an approved integrated revision and preserved the original. "
            "Reassess the output without supplemental answers."
        ),
        mode="integrated",
        plan_digest=approved_plan.plan.plan_digest,
        final_assessment_required=True,
    )


def _write_docx_revision(
    source: Path, output: Path, grouped: dict[str, list[str]]
) -> None:
    document = Document(source)
    document.add_page_break()
    document.add_heading("Forge Clarifications", level=1)
    document.add_paragraph(
        "User-provided clarifications added after the initial readiness review."
    )
    for criterion_id, answers in grouped.items():
        document.add_heading(_criterion_title(criterion_id), level=2)
        for answer in answers:
            document.add_paragraph(answer)
    document.save(output)


def _write_text_revision(
    source: Path,
    output: Path,
    grouped: dict[str, list[str]],
    *,
    markdown: bool,
) -> None:
    heading = "# Forge Clarifications" if markdown else "FORGE CLARIFICATIONS"
    rendered = [
        source.read_text(encoding="utf-8").rstrip(),
        "",
        heading,
        "",
        "User-provided clarifications added after the initial readiness review.",
    ]
    for criterion_id, answers in grouped.items():
        title = _criterion_title(criterion_id)
        rendered.extend(["", f"## {title}" if markdown else title.upper(), ""])
        rendered.extend(answers)
    output.write_text("\n".join(rendered).rstrip() + "\n", encoding="utf-8")


def _criterion_title(criterion_id: str) -> str:
    return criterion_id.replace("_", " ").title()


def _validate_source(
    source: Path, supplemental_answers: list[SupplementalAnswer]
) -> None:
    if not source.exists() or not source.is_file():
        raise FileNotFoundError(f"document not found: {source}")
    if not supplemental_answers:
        raise ValueError("at least one supplemental answer is required")
    if source.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"unsupported document type {source.suffix.lower()!r}")
    if source.suffix.lower() == ".pdf":
        raise ValueError("PDF revisions require an editable source")


def _validate_paths(
    source: Path,
    output: Path,
    supplemental_answers: list[SupplementalAnswer],
) -> None:
    _validate_source(source, supplemental_answers)
    if source == output:
        raise ValueError("output_path must differ from source_path")
    if output.exists():
        raise ValueError(f"output already exists: {output}")
    if not output.parent.exists():
        raise FileNotFoundError(f"output directory not found: {output.parent}")
    if output.suffix.lower() != source.suffix.lower():
        raise ValueError("output_path must use the same file type as source_path")


def _file_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _canonical_digest(payload: dict[str, object]) -> str:
    encoded = json.dumps(
        payload, sort_keys=True, separators=(",", ":"), ensure_ascii=False
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _verify_plan_digest(plan: RevisionPlan) -> None:
    payload = plan.model_dump(exclude={"plan_digest"})
    if _canonical_digest(payload) != plan.plan_digest:
        raise ValueError("revision plan was modified after preview")


def _source_sections(source: Path) -> list[str]:
    suffix = source.suffix.lower()
    if suffix == ".docx":
        return [
            paragraph.text.strip()
            for paragraph in Document(source).paragraphs
            if paragraph.text.strip()
            and (paragraph.style.name or "").lower().startswith("heading")
        ]
    if suffix == ".md":
        return [
            match.group(1).strip()
            for line in source.read_text(encoding="utf-8").splitlines()
            if (match := re.match(r"^#{1,6}\s+(.+?)\s*$", line))
        ]
    return []


def _source_text(source: Path) -> str:
    if source.suffix.lower() == ".docx":
        return "\n".join(paragraph.text for paragraph in Document(source).paragraphs)
    return source.read_text(encoding="utf-8")


def _best_section(criterion_id: str, sections: list[str]) -> str | None:
    aliases = _SECTION_ALIASES.get(
        criterion_id, (criterion_id.replace("_", " "),)
    )
    ranked: list[tuple[int, int, str]] = []
    for index, section in enumerate(sections):
        normalized = section.casefold()
        score = sum(1 for alias in aliases if alias in normalized)
        if score:
            ranked.append((-score, index, section))
    return min(ranked)[2] if ranked else None


def _section_excerpt(source: Path, target: str) -> str | None:
    if source.suffix.lower() == ".docx":
        paragraphs = list(Document(source).paragraphs)
        heading_index = next(
            (
                index
                for index, paragraph in enumerate(paragraphs)
                if paragraph.text.strip().casefold() == target.casefold()
                and (paragraph.style.name or "").lower().startswith("heading")
            ),
            None,
        )
        if heading_index is None:
            return None
        body: list[str] = []
        for paragraph in paragraphs[heading_index + 1 :]:
            if (paragraph.style.name or "").lower().startswith("heading"):
                break
            if paragraph.text.strip():
                body.append(paragraph.text.strip())
        rendered = " ".join(body).strip()
        return rendered[:500] or None
    if source.suffix.lower() == ".md":
        text = source.read_text(encoding="utf-8")
        pattern = re.compile(
            rf"^#{{1,6}}\s+{re.escape(target)}\s*$",
            re.MULTILINE | re.IGNORECASE,
        )
        match = pattern.search(text)
        if match is None:
            return None
        next_heading = re.search(r"^#{1,6}\s+", text[match.end() :], re.MULTILINE)
        end = match.end() + next_heading.start() if next_heading else len(text)
        rendered = text[match.end() : end].strip()
        return rendered[:500] or None
    return None


def _write_integrated_markdown(
    source: Path,
    output: Path,
    edits: list[ApprovedRevisionEdit],
    plan: ApprovedRevisionPlan,
) -> None:
    text = source.read_text(encoding="utf-8")
    insertions: list[tuple[int, str]] = []
    for approved in edits:
        heading = approved.edit.target_section
        pattern = re.compile(
            rf"^#{{1,6}}\s+{re.escape(heading or '')}\s*$",
            re.MULTILINE | re.IGNORECASE,
        )
        match = pattern.search(text)
        if match is None:
            raise ValueError(f"revision target section {heading!r} no longer exists")
        next_heading = re.search(r"^#{1,6}\s+", text[match.end() :], re.MULTILINE)
        position = (
            match.end() + next_heading.start()
            if next_heading is not None
            else len(text)
        )
        insertions.append((position, f"\n\n{approved.edit.answer}\n"))
    for position, content in sorted(insertions, reverse=True):
        text = text[:position].rstrip() + content + text[position:].lstrip("\n")
    text = text.rstrip() + _render_audit_markdown(plan) + "\n"
    output.write_text(text, encoding="utf-8")


def _write_integrated_text(
    source: Path,
    output: Path,
    edits: list[ApprovedRevisionEdit],
    plan: ApprovedRevisionPlan,
) -> None:
    if edits:
        raise ValueError("plain text has no reliable section targets; use audit_only")
    text = source.read_text(encoding="utf-8").rstrip()
    output.write_text(text + _render_audit_text(plan) + "\n", encoding="utf-8")


def _write_integrated_docx(
    source: Path,
    output: Path,
    edits: list[ApprovedRevisionEdit],
    plan: ApprovedRevisionPlan,
) -> None:
    document = Document(source)
    for approved in edits:
        target = approved.edit.target_section
        paragraphs = list(document.paragraphs)
        heading_index = next(
            (
                index
                for index, paragraph in enumerate(paragraphs)
                if paragraph.text.strip().casefold() == (target or "").casefold()
                and (paragraph.style.name or "").lower().startswith("heading")
            ),
            None,
        )
        if heading_index is None:
            raise ValueError(f"revision target section {target!r} no longer exists")
        anchor: Paragraph = paragraphs[heading_index]
        for paragraph in paragraphs[heading_index + 1 :]:
            if (paragraph.style.name or "").lower().startswith("heading"):
                break
            anchor = paragraph
        inserted = document.add_paragraph(approved.edit.answer)
        anchor._p.addnext(inserted._p)
    document.add_page_break()
    document.add_heading("Forge Revision Audit", level=1)
    document.add_paragraph(f"Approved plan: {plan.approval_digest}")
    for approved in plan.edits:
        if approved.action == "skip":
            continue
        document.add_heading(_criterion_title(approved.edit.criterion_id), level=2)
        document.add_paragraph(f"Action: {approved.action}")
        document.add_paragraph(approved.edit.answer)
    document.save(output)


def _render_audit_markdown(plan: ApprovedRevisionPlan) -> str:
    lines = ["", "# Forge Revision Audit", "", f"Approved plan: `{plan.approval_digest}`"]
    for approved in plan.edits:
        if approved.action == "skip":
            continue
        lines.extend(
            [
                "",
                f"## {_criterion_title(approved.edit.criterion_id)}",
                "",
                f"Action: {approved.action}",
                "",
                approved.edit.answer,
            ]
        )
    return "\n".join(lines)


def _render_audit_text(plan: ApprovedRevisionPlan) -> str:
    lines = ["", "", "FORGE REVISION AUDIT", f"Approved plan: {plan.approval_digest}"]
    for approved in plan.edits:
        if approved.action == "skip":
            continue
        lines.extend(
            [
                "",
                _criterion_title(approved.edit.criterion_id).upper(),
                f"Action: {approved.action}",
                approved.edit.answer,
            ]
        )
    return "\n".join(lines)
